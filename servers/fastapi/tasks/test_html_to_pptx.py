"""HTML → PPTX (deckhtml) compile tests."""

from __future__ import annotations

import os
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import Mock, patch


class HtmlToPptxCompileTests(unittest.TestCase):
    def test_failsafe_writer_emits_only_editable_shapes(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        from tasks.html_to_pptx import build_failsafe_pptx

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        html_dir.mkdir()
        (html_dir / "slide_01.html").write_text(
            "<html><body><h1>Emergency title</h1><p>Preserved body copy.</p></body></html>",
            encoding="utf-8",
        )
        out = root / "fallback.pptx"
        meta = build_failsafe_pptx(root, out)

        self.assertTrue(out.is_file())
        self.assertEqual("editable_text_failsafe", meta["mode"])
        self.assertEqual(1, meta["editable_slides"])
        self.assertEqual(0, meta["raster_slides"])
        prs = Presentation(str(out))
        self.assertEqual(1, len(prs.slides))
        self.assertFalse(
            any(
                shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                for shape in prs.slides[0].shapes
            )
        )
        text = " ".join(
            shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text")
        )
        self.assertIn("Emergency title", text)
        self.assertIn("Preserved body copy", text)

    def test_finalize_compile_writes_done_not_completed(self):
        from tasks.generator_agent import finalize_compile

        root = Path(tempfile.mkdtemp())
        (root / "html_output").mkdir()
        (root / "html_output" / "slide_01.html").write_text(
            "<!DOCTYPE html><html><body><section class='slide'></section></body></html>\n",
            encoding="utf-8",
        )
        written: list[tuple] = []

        def _fake_write(storage_dir, status, progress=0, **kwargs):
            written.append((status, progress, kwargs.get("message_key")))

        generate_task = types.ModuleType("tasks.generate_task")
        generate_task.write_status = _fake_write
        generate_task.send_ws_message = lambda *_args, **_kwargs: None

        with (
            patch.dict(sys.modules, {"tasks.generate_task": generate_task}),
            patch(
                "tasks.html_to_pptx.finalize_compile_html",
                return_value={
                    "ok": True,
                    "release_manifest": {"status": "released"},
                },
            ),
            patch("asyncio.new_event_loop") as loop_factory,
        ):
            loop = loop_factory.return_value
            loop.run_until_complete.return_value = None
            result = finalize_compile(str(root), project_id="p1", run_id="r1")

        self.assertTrue(result.get("ok"))
        terminal = [w for w in written if w[1] == 100]
        self.assertTrue(terminal)
        self.assertEqual(terminal[-1][0], "done")

    def test_compile_html_dir_invokes_deckhtml(self):
        from tasks.html_to_pptx import compile_html_dir_to_pptx

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        html_dir.mkdir()
        (html_dir / "slide_01.html").write_text(
            "<html><body><section class='slide'><h1>Hi</h1></section></body></html>\n",
            encoding="utf-8",
        )
        out = root / "out.pptx"
        seen: list[list[str]] = []

        def _fake_run(cmd, **kwargs):
            seen.append(list(cmd))
            Path(cmd[cmd.index("-o") + 1]).write_bytes(b"PK\x03\x04fake")

            class R:
                returncode = 0
                stdout = "ok"
                stderr = ""

            return R()

        with patch("tasks.html_to_pptx.subprocess.run", side_effect=_fake_run):
            path = compile_html_dir_to_pptx(html_dir, out)
        self.assertEqual(Path(path), out)
        self.assertTrue(out.is_file())
        self.assertTrue(seen)
        # Prefer vendored deckhtml (node …/vendor/deckhtml/dist/cli.js) or tasks/.bin
        joined = " ".join(seen[0])
        self.assertTrue(
            "vendor/deckhtml" in joined or "node_modules" in joined or seen[0][0].endswith("deckhtml"),
            msg=joined,
        )
        self.assertIn("--no-slide-transitions", seen[0])

    def test_compile_passes_fixed_slide_transition(self):
        from tasks.html_to_pptx import compile_html_dir_to_pptx

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        html_dir.mkdir()
        (html_dir / "slide_01.html").write_text(
            "<html><body><section class='slide'><h1>Hi</h1></section></body></html>\n",
            encoding="utf-8",
        )
        out = root / "out.pptx"
        seen: list[list[str]] = []

        def _fake_run(cmd, **kwargs):
            seen.append(list(cmd))
            Path(cmd[cmd.index("-o") + 1]).write_bytes(b"PK\x03\x04fake")

            class R:
                returncode = 0
                stdout = "ok"
                stderr = ""

            return R()

        with patch("tasks.html_to_pptx.subprocess.run", side_effect=_fake_run):
            compile_html_dir_to_pptx(html_dir, out, slide_transition="fade")
        self.assertIn("--slide-transition", seen[0])
        self.assertEqual(seen[0][seen[0].index("--slide-transition") + 1], "fade")
        self.assertNotIn("--no-slide-transitions", seen[0])

    def test_deckhtml_bin_resolves_vendored_package(self):
        from tasks import html_to_pptx as mod

        tasks_dir = Path(mod.__file__).resolve().parent
        vendor_pkg = tasks_dir.parent / "vendor" / "deckhtml" / "package.json"
        self.assertTrue(vendor_pkg.is_file(), "vendored deckhtml missing")
        with patch.dict(os.environ, {}, clear=False):
            os.environ.pop("DECKHTML_BIN", None)
            bin_argv = mod._deckhtml_bin()
        joined = " ".join(bin_argv)
        # Prefer tasks/node_modules/.bin (file: → vendor) or direct vendor/dist/cli.js
        resolved = Path(bin_argv[-1]).resolve()
        self.assertTrue(
            "vendor/deckhtml" in joined
            or "vendor/deckhtml" in str(resolved)
            or (tasks_dir / "node_modules" / "@deckflow" / "deckhtml").is_dir(),
            msg=joined,
        )

    def test_list_html_slides_sorted(self):
        from tasks.html_to_pptx import list_html_slides

        root = Path(tempfile.mkdtemp())
        (root / "slide_02.html").write_text("b", encoding="utf-8")
        (root / "slide_01.html").write_text("a", encoding="utf-8")
        names = [p.name for p in list_html_slides(root)]
        self.assertEqual(names, ["slide_01.html", "slide_02.html"])

    def test_embed_speaker_notes_matches_html_order(self):
        from pptx import Presentation
        from pptx.util import Inches

        from tasks.html_to_pptx import embed_speaker_notes_into_pptx

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        notes_dir = root / "notes"
        html_dir.mkdir()
        notes_dir.mkdir()
        (html_dir / "slide_01.html").write_text("<html></html>\n", encoding="utf-8")
        (html_dir / "slide_02.html").write_text("<html></html>\n", encoding="utf-8")
        (notes_dir / "slide_01.md").write_text(
            "Good morning. This is the cover spoken note with enough words.\n",
            encoding="utf-8",
        )
        (notes_dir / "slide_02.md").write_text(
            "Next we cover the method and why repair becomes the bottleneck in practice.\n",
            encoding="utf-8",
        )

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
        prs.slides.add_slide(blank)
        prs.slides.add_slide(blank)
        pptx = root / "deck.pptx"
        prs.save(str(pptx))

        meta = embed_speaker_notes_into_pptx(root, pptx)
        self.assertEqual(meta["embedded"], 2)
        loaded = Presentation(str(pptx))
        self.assertIn("cover spoken note", loaded.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn("repair becomes", loaded.slides[1].notes_slide.notes_text_frame.text)

    def test_finalize_skips_notes_when_disabled(self):
        from tasks.html_to_pptx import finalize_compile_html

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        html_dir.mkdir()
        (html_dir / "slide_01.html").write_text("<html></html>\n", encoding="utf-8")
        (root / "exports").mkdir()

        embed = Mock()
        with patch(
            "tasks.html_to_pptx.validate_html_dir_runtime_layout",
            return_value={"ok": True, "hard": [], "slides": []},
        ), patch(
            "tasks.html_to_pptx.compile_html_dir_to_pptx",
            side_effect=lambda _html, output, **_k: Path(output).write_bytes(b"PK\x03\x04fake") or str(output),
        ), patch(
            "tasks.html_to_pptx.embed_speaker_notes_into_pptx", new=embed
        ), patch(
            "tasks.release_export.write_release_manifest",
            return_value={"status": "released"},
        ):
            result = finalize_compile_html(
                str(root), include_notes=False, fidelity_verify=False
            )
        embed.assert_not_called()
        self.assertEqual(result["speaker_notes"].get("reason"), "disabled")
        self.assertIn("exports/releases/", result["pptx_path"])

    def test_finalize_native_failure_without_previous_release_does_not_publish_plain_deck(self):
        from tasks.html_to_pptx import HtmlToPptxError, finalize_compile_html

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        html_dir.mkdir()
        (html_dir / "slide_01.html").write_text("<html><h1>Fallback</h1></html>")

        publish = Mock()
        with patch(
            "tasks.html_to_pptx.validate_html_dir_runtime_layout",
            return_value={"ok": True, "hard": [], "slides": []},
        ), patch(
            "tasks.html_to_pptx.compile_html_dir_to_pptx",
            side_effect=HtmlToPptxError("browser unavailable"),
        ), patch("tasks.release_export.write_release_manifest", new=publish):
            with self.assertRaisesRegex(HtmlToPptxError, "previous immutable release"):
                finalize_compile_html(
                    str(root), include_notes=False, fidelity_verify=False
                )
        publish.assert_not_called()

    def test_finalize_native_failure_serves_explicit_previous_release(self):
        from tasks.html_to_pptx import HtmlToPptxError, finalize_compile_html

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        html_dir.mkdir()
        (html_dir / "slide_01.html").write_text("<html><h1>Draft</h1></html>")
        previous = root / "exports" / "releases" / "prior" / "presentation.pptx"
        previous.parent.mkdir(parents=True)
        previous.write_bytes(b"PK\x03\x04prior")
        previous_manifest = {
            "status": "released",
            "pptx_path": "exports/releases/prior/presentation.pptx",
        }

        with patch(
            "tasks.html_to_pptx.validate_html_dir_runtime_layout",
            return_value={"ok": True, "hard": [], "slides": []},
        ), patch(
            "tasks.html_to_pptx.compile_html_dir_to_pptx",
            side_effect=HtmlToPptxError("browser unavailable"),
        ), patch(
            "tasks.release_export.resolve_previous_released_pptx",
            return_value=(str(previous), "sha", previous_manifest),
        ):
            result = finalize_compile_html(
                str(root), include_notes=False, fidelity_verify=False
            )
        self.assertTrue(result["ok"])
        self.assertEqual("previous_immutable_release", result["mode"])
        self.assertFalse(result["fallback"]["whole_slide_replaced"])

    def test_fidelity_failure_records_object_remediation_without_replacing_slide(self):
        from tasks.html_to_pptx import finalize_compile_html

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        preview_dir = root / "preview"
        html_dir.mkdir()
        preview_dir.mkdir()
        (html_dir / "slide_01.html").write_text("<html><h1>Native</h1></html>")
        (preview_dir / "slide_01.png").write_bytes(b"png")

        def compile_staged(_html, output, **_kwargs):
            Path(output).write_bytes(b"PK\x03\x04native")
            return str(output)

        fallback = Mock()
        with patch(
            "tasks.html_to_pptx.validate_html_dir_runtime_layout",
            return_value={"ok": True, "hard": [], "slides": []},
        ), patch(
            "tasks.html_to_pptx.compile_html_dir_to_pptx",
            side_effect=compile_staged,
        ), patch(
            "tasks.pptx_visual_verify.verify_pptx_against_previews",
            return_value={"ok": False, "failures": [{"stem": "slide_01"}]},
        ), patch(
            "tasks.html_to_pptx.build_failsafe_pptx", new=fallback
        ), patch(
            "tasks.release_export.write_release_manifest",
            return_value={"status": "released_with_advisories"},
        ):
            result = finalize_compile_html(
                str(root), include_notes=False, fidelity_verify=True
            )

        fallback.assert_not_called()
        self.assertTrue(result["ok"])
        self.assertEqual("html_deckhtml", result["mode"])
        self.assertIn(
            "object_remediation_recorded_nonblocking",
            [item.get("code") for item in result.get("advisories") or []],
        )
        plan_path = root / "validation" / "native_fidelity_remediation.json"
        self.assertTrue(plan_path.is_file())
        plan = __import__("json").loads(plan_path.read_text(encoding="utf-8"))
        self.assertFalse(plan["whole_slide_fallback_allowed"])

    def test_finalize_note_failure_preserves_native_deck_and_source_notes(self):
        from tasks.html_to_pptx import finalize_compile_html

        root = Path(tempfile.mkdtemp())
        html_dir = root / "html_output"
        html_dir.mkdir()
        (html_dir / "slide_01.html").write_text("<html><h1>Notes</h1></html>")

        def compile_staged(_html, output, **_kwargs):
            Path(output).write_bytes(b"PK\x03\x04fake")
            return str(output)

        with patch(
            "tasks.html_to_pptx.validate_html_dir_runtime_layout",
            return_value={"ok": True, "hard": [], "slides": []},
        ), patch(
            "tasks.html_to_pptx.compile_html_dir_to_pptx", side_effect=compile_staged
        ), patch(
            "tasks.html_to_pptx.embed_speaker_notes_into_pptx",
            side_effect=RuntimeError("notes XML failure"),
        ), patch(
            "tasks.release_export.write_release_manifest",
            return_value={"status": "released_with_advisories"},
        ):
            result = finalize_compile_html(
                str(root), include_notes=True, fidelity_verify=False
            )
        self.assertTrue(result["ok"])
        self.assertTrue(result["speaker_notes"]["preserved_source_notes"])
        self.assertIn(
            "embedding_failed_nonblocking",
            [item.get("code") for item in result.get("advisories") or []],
        )


if __name__ == "__main__":
    unittest.main()
