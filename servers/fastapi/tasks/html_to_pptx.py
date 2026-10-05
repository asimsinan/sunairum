"""Compile html_output/ → PPTX via @deckflow/deckhtml (local mode).

HTML is the product source of truth. Slide IR / ir_to_pptx is retired.
Intelligent speaker notes from notes/slide_*.md are merged into the PPTX
after deckhtml (python-pptx), matching slide order to html_output.
"""

from __future__ import annotations

import json
import html as html_lib
import os
import re
import shutil
import subprocess
import uuid
from pathlib import Path
from typing import Any


class HtmlToPptxError(RuntimeError):
    """deckhtml conversion failed."""


def _deckhtml_bin() -> list[str]:
    """Resolve deckhtml CLI: DECKHTML_BIN, tasks/node_modules, or vendored dist."""
    override = (os.environ.get("DECKHTML_BIN") or "").strip()
    if override:
        return override.split()
    here = Path(__file__).resolve().parent
    local = here / "node_modules" / ".bin" / "deckhtml"
    if local.is_file():
        return [str(local)]
    vendor_root = here.parent / "vendor" / "deckhtml"
    vendor_cli = vendor_root / "dist" / "cli.js"
    if vendor_cli.is_file() or (vendor_root / "package.json").is_file():
        return ["node", str(vendor_cli)]
    root_nm = here.parent / "node_modules" / ".bin" / "deckhtml"
    if root_nm.is_file():
        return [str(root_nm)]
    which = shutil.which("deckhtml")
    if which:
        return [which]
    raise HtmlToPptxError(
        "deckhtml not found; build servers/fastapi/vendor/deckhtml "
        "and npm install in servers/fastapi/tasks"
    )


def list_html_slides(html_dir: str | Path) -> list[Path]:
    root = Path(html_dir)
    files = sorted(root.glob("slide_*.html"))
    return [p for p in files if p.is_file() and p.stat().st_size > 0]


def validate_html_dir_runtime_layout(
    storage_dir: str | Path,
    html_dir: str | Path,
) -> dict[str, Any]:
    """Measure every source HTML slide and persist a release-audit report.

    This is deliberately performed immediately before compilation so editor
    changes cannot bypass the validation performed during generation.
    """
    from tasks.html_layout_qa import audit_runtime_html_layout
    from tasks.html_pipeline import (
        inject_formula_into_html,
        inject_pdf_figure_into_html,
        inject_table_into_html,
        normalize_html_font_families,
        relax_frozen_text_heights,
        normalize_inline_text_boundaries,
        normalize_figure_hosts,
        sanitize_html_audience_copy,
    )
    from tasks.scene_html_author import normalize_scene_authored_html, validate_scene_authored_html
    from tasks.scene_local_repair import repair_scene_html_until_stable

    reports: list[dict[str, Any]] = []
    hard: list[str] = []
    storage_root = Path(storage_dir)
    slide_payloads: dict[int, dict[str, Any]] = {}
    strategist_path = storage_root / "validation" / "strategist_ir.json"
    if strategist_path.is_file():
        try:
            strategist = json.loads(strategist_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            strategist = {}
        for raw in strategist.get("slides") or []:
            if not isinstance(raw, dict):
                continue
            try:
                slide_payloads[int(raw.get("slide") or 0)] = raw
            except (TypeError, ValueError):
                continue
    for html_path in list_html_slides(html_dir):
        scene_path = storage_root / "scene" / f"{html_path.stem}.json"
        scene_payload: dict[str, Any] = {}
        scene_valid = False
        if scene_path.is_file():
            try:
                scene_payload = json.loads(scene_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                scene_payload = {}
            if scene_payload.get("version") == "scene_layout_v2":
                scene_valid = True
        local_repair: dict[str, Any] = {"changed": False, "regions": []}
        structural: list[str] = []
        try:
            html = html_path.read_text(encoding="utf-8")
        except OSError as exc:
            report = {"ok": False, "issues": [f"html_read_failed: {exc}"], "hard": [f"html_read_failed: {exc}"], "inspection": {}}
        else:
            html, audience_copy = sanitize_html_audience_copy(html)
            if audience_copy.get("changed"):
                html_path.write_text(html if html.endswith("\n") else html + "\n", encoding="utf-8")
            html, inline_boundaries = normalize_inline_text_boundaries(html)
            if inline_boundaries.get("changed"):
                html_path.write_text(html if html.endswith("\n") else html + "\n", encoding="utf-8")
            # A deck may only name faces the exporter can embed; anything else
            # resolves differently in Chromium, LibreOffice and the reader's
            # PowerPoint, and the text drifts out of its measured box.
            html, font_families = normalize_html_font_families(html)
            html, frozen_heights = relax_frozen_text_heights(html)
            if frozen_heights.get("changed"):
                html_path.write_text(html if html.endswith("\n") else html + "\n", encoding="utf-8")
            if font_families.get("changed"):
                html_path.write_text(html if html.endswith("\n") else html + "\n", encoding="utf-8")
            if scene_valid:
                html, scene_normalization = normalize_scene_authored_html(html, scene_payload)
                if scene_normalization.get("changed"):
                    html_path.write_text(html if html.endswith("\n") else html + "\n", encoding="utf-8")
            # Normalize semantic-host ownership before the Scene constitution
            # check so rollout migration can remove a legacy second <style>
            # tag instead of being blocked by the artifact it is designed to
            # repair.
            html, early_figure_chrome = normalize_figure_hosts(html)
            if early_figure_chrome.get("changed"):
                html_path.write_text(html if html.endswith("\n") else html + "\n", encoding="utf-8")
            structural = validate_scene_authored_html(html, scene_payload) if scene_valid else []
            try:
                slide_number = int(html_path.stem.rsplit("_", 1)[-1])
            except (TypeError, ValueError):
                slide_number = 0
            slide_payload = dict(slide_payloads.get(slide_number) or {})
            slide_payload.setdefault("slide", slide_number)
            # Semantic hosts are pipeline-owned. Refresh them deterministically
            # after editor changes or legacy fallback migration; no visual LLM
            # is called and authored surrounding composition is preserved.
            if not structural:
                empty_figure = re.search(
                    r'<div\b[^>]*data-pptx=["\'](?:figure|image)["\'][^>]*>\s*</div>',
                    html,
                    flags=re.I | re.S,
                )
                figure_id = str(slide_payload.get("figure_id") or "").strip()
                if empty_figure and figure_id:
                    html, _figure_meta = inject_pdf_figure_into_html(
                        html, storage_dir=str(storage_root), figure_id=figure_id
                    )
                html, _figure_chrome = normalize_figure_hosts(html)

                pre_host_report = audit_runtime_html_layout(html)
                table_box = None
                for element in (pre_host_report.get("inspection") or {}).get("elements") or []:
                    if str(element.get("pptx") or "").casefold() == "table":
                        table_box = element.get("box")
                        break
                if table_box:
                    html, _table_meta = inject_table_into_html(
                        html,
                        storage_dir=str(storage_root),
                        slide=slide_payload,
                        host_box=table_box,
                    )
                if re.search(
                    r'<div\b[^>]*data-pptx=["\']formula["\'][^>]*>\s*</div>',
                    html,
                    flags=re.I | re.S,
                ):
                    html, _formula_meta = inject_formula_into_html(
                        html, storage_dir=str(storage_root), slide=slide_payload
                    )
                html_path.write_text(html if html.endswith("\n") else html + "\n", encoding="utf-8")
            report = audit_runtime_html_layout(html)
            inspection = report.get("inspection") if isinstance(report.get("inspection"), dict) else {}
            # Editor changes and older generated slides take the same bounded,
            # measured repair path. Never call an LLM or replace the composition.
            if not structural and inspection.get("available"):
                html, report, local_repair = repair_scene_html_until_stable(
                    html,
                    scene_payload,
                    initial_report=report,
                )
                if local_repair.get("changed"):
                    html_path.write_text(
                        html if html.endswith("\n") else html + "\n",
                        encoding="utf-8",
                    )
            # Fidelity references must be rendered from the exact canonical
            # HTML revision being compiled. An older validator may already
            # have persisted a repair without refreshing its PNG, so checking
            # only this call's string mutation can retain stale broken images.
            if not report.get("hard"):
                from tasks.html_pipeline import _render_preview_png

                preview_path = storage_root / "preview" / f"{html_path.stem}.png"
                preview_path.parent.mkdir(parents=True, exist_ok=True)
                if not _render_preview_png(html, str(preview_path)):
                    report.setdefault("issues", []).append("preview_refresh_failed")
                    report.setdefault("hard", []).append("preview_refresh_failed")
        slide_hard = [str(issue) for issue in report.get("hard") or []]
        slide_hard.extend(f"scene_contract:{item}" for item in structural)
        if not scene_valid:
            slide_hard.insert(0, "scene_manifest_missing_or_invalid")
        hard.extend(f"{html_path.stem}: {issue}" for issue in slide_hard)
        from tasks.layout_findings import classify_runtime_findings

        reports.append(
            {
                "stem": html_path.stem,
                "ok": not slide_hard,
                "issues": [str(issue) for issue in report.get("issues") or []],
                "hard": slide_hard,
                "findings": classify_runtime_findings(slide_hard),
                "inspection": report.get("inspection") or {},
                "layout_generation": "scene_layout_v2" if scene_valid else "invalid",
                "scene_contract": structural,
                "local_layout_repair": local_repair,
            }
        )
    result = {
        "ok": not hard,
        "hard": hard,
        "findings": [
            finding
            for slide_report in reports
            for finding in (slide_report.get("findings") or [])
        ],
        "slides": reports,
        "summary": {
            "scene_layout_v2": sum(
                1 for item in reports if item["layout_generation"] == "scene_layout_v2"
            ),
            "invalid_scene_manifest": sum(
                1 for item in reports if item["layout_generation"] == "invalid"
            ),
        },
    }
    audit_path = Path(storage_dir) / "validation" / "runtime_layout_qa.json"
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = audit_path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temporary, audit_path)
    return result


def compile_html_dir_to_pptx(
    html_dir: str | Path,
    output_path: str | Path,
    *,
    timeout: int | None = None,
    slide_transition: str | None = None,
    embed_fonts: bool = True,
) -> str:
    """Run deckhtml --mode local on slide_*.html → output_path.

    ``slide_transition`` is a catalog name (same effect on every slide) or
    ``random`` (different effect per slide). ``None`` disables transitions.
    Local mode embeds deck fonts by default (``embed_fonts=True``).
    """
    slides = list_html_slides(html_dir)
    if not slides:
        raise FileNotFoundError(f"No slide_*.html in {html_dir}")
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        *_deckhtml_bin(),
        *[str(p) for p in slides],
        "--mode",
        "local",
        "--force",
        "--identity-attribute",
        "data-layout-id",
        "--report",
        "-o",
        str(out),
    ]
    if embed_fonts:
        cmd.append("--embed-fonts")
    token = str(slide_transition or "").strip()
    if token:
        cmd.extend(["--slide-transition", token])
    else:
        cmd.append("--no-slide-transitions")
    env = os.environ.copy()
    # Prefer local conversion; never prompt for cloud auth in workers.
    env.setdefault("CI", "1")
    # Point deck-font-embedder at curated TTFs shipped with the FastAPI image.
    fonts_dir = Path(__file__).resolve().parents[1] / "assets" / "fonts"
    if fonts_dir.is_dir():
        env.setdefault("DECKHTML_FONTS_DIR", str(fonts_dir))
    limit = timeout or int(os.environ.get("DECKHTML_TIMEOUT_SECONDS") or "600")
    try:
        proc = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=limit,
            env=env,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise HtmlToPptxError(f"deckhtml timed out after {limit}s") from exc
    except FileNotFoundError as exc:
        raise HtmlToPptxError(
            "deckhtml not found; install @deckflow/deckhtml in the worker image"
        ) from exc
    if proc.returncode != 0 or not out.is_file() or out.stat().st_size == 0:
        stderr = (proc.stderr or "").strip()
        stdout = (proc.stdout or "").strip()
        detail = stderr or stdout or f"exit {proc.returncode}"
        raise HtmlToPptxError(f"deckhtml failed: {detail[:2000]}")
    return str(out)


_SLIDE_STEM_RE = re.compile(r"^(slide_\d+)$", re.I)


def _spoken_note_for_stem(storage_dir: str | Path, stem: str) -> str:
    """Load notes/{stem}.md and clean to spoken plain prose."""
    from tasks.speaker_notes import (
        _clean_spoken_note,
        dedupe_spoken_note,
        note_is_usable,
        read_note,
    )

    path = Path(storage_dir) / "notes" / f"{stem}.md"
    raw = read_note(path)
    if not note_is_usable(raw):
        return ""
    cleaned = _clean_spoken_note(raw)
    return dedupe_spoken_note(cleaned)


def embed_speaker_notes_into_pptx(
    storage_dir: str | Path,
    pptx_path: str | Path,
    *,
    html_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Merge notes/slide_*.md into PPTX speaker notes (deckhtml has no notes path).

    Matching is intelligent:
    - PPTX slide i ↔ ith sorted html_output/slide_*.html stem
    - Skip empty / unusable notes; strip markdown chrome; dedupe spoken repeats
    - Never invent notes when the markdown file is missing
    """
    from pptx import Presentation

    root = Path(storage_dir)
    pptx = Path(pptx_path)
    slides_html = list_html_slides(html_dir or (root / "html_output"))
    if not slides_html:
        return {"embedded": 0, "skipped": 0, "reason": "no_html_slides"}
    if not pptx.is_file():
        raise FileNotFoundError(f"PPTX not found: {pptx}")

    prs = Presentation(str(pptx))
    pptx_count = len(prs.slides)
    embedded = 0
    skipped = 0
    for index, html_path in enumerate(slides_html):
        if index >= pptx_count:
            skipped += len(slides_html) - index
            break
        stem = html_path.stem
        if not _SLIDE_STEM_RE.match(stem):
            skipped += 1
            continue
        note = _spoken_note_for_stem(root, stem)
        if not note:
            skipped += 1
            continue
        slide = prs.slides[index]
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = note
        embedded += 1

    if embedded:
        prs.save(str(pptx))
    return {
        "embedded": embedded,
        "skipped": skipped,
        "pptx_slides": pptx_count,
        "html_slides": len(slides_html),
    }


def _plain_html_text(document: str) -> str:
    """Extract readable fallback copy without treating authored HTML as instructions."""
    text = re.sub(r"<script\b[\s\S]*?</script>", " ", document or "", flags=re.I)
    text = re.sub(r"<style\b[\s\S]*?</style>", " ", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", html_lib.unescape(text)).strip()


def _html_heading(document: str, fallback: str) -> str:
    match = re.search(r"<h1\b[^>]*>([\s\S]*?)</h1>", document or "", flags=re.I)
    if not match:
        match = re.search(r"<title\b[^>]*>([\s\S]*?)</title>", document or "", flags=re.I)
    heading = _plain_html_text(match.group(1)) if match else ""
    return heading or fallback


def build_failsafe_pptx(
    storage_dir: str | Path,
    output_path: str | Path,
    *,
    html_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Always produce an editable 16:9 PPTX without DeckHTML.

    This is the last-resort availability writer, not a fidelity substitute for
    a successful native compilation. Every fallback slide is made from native
    PowerPoint text boxes and shapes. Full-slide screenshots are deliberately
    forbidden here: a downloadable deck must never silently become an image
    deck merely because the native compiler or fidelity verifier had a problem.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    root = Path(storage_dir)
    source_dir = Path(html_dir or (root / "html_output"))
    slides = list_html_slides(source_dir)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    editable_slides = 0
    errors: list[str] = []

    if not slides:
        slides = [source_dir / "slide_01.html"]

    for index, html_path in enumerate(slides, start=1):
        try:
            document = html_path.read_text(encoding="utf-8")
        except OSError as exc:
            document = ""
            errors.append(f"{html_path.stem}: html_read_failed: {exc}")

        slide = prs.slides.add_slide(blank)
        editable_slides += 1
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(248, 250, 252)
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(0.45), Inches(0.06), Inches(4.72)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(47, 132, 214)
        bar.line.fill.background()
        title = _html_heading(document, f"Slide {index}")
        body = _plain_html_text(document)
        if body.startswith(title):
            body = body[len(title) :].strip()
        title_box = slide.shapes.add_textbox(
            Inches(0.72), Inches(0.62), Inches(8.7), Inches(1.2)
        )
        title_frame = title_box.text_frame
        title_frame.clear()
        title_frame.word_wrap = True
        run = title_frame.paragraphs[0].add_run()
        run.text = title[:240]
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor(26, 54, 93)
        body_box = slide.shapes.add_textbox(
            Inches(0.74), Inches(2.0), Inches(8.55), Inches(2.75)
        )
        body_frame = body_box.text_frame
        body_frame.clear()
        body_frame.word_wrap = True
        body_run = body_frame.paragraphs[0].add_run()
        body_run.text = body[:2200] or "This slide could not be rendered from HTML."
        body_run.font.name = "Arial"
        body_run.font.size = Pt(16)
        body_run.font.color.rgb = RGBColor(71, 85, 105)

    prs.save(str(out))
    if not out.is_file() or out.stat().st_size == 0:
        raise HtmlToPptxError("Fail-safe PPTX writer produced no artifact")
    return {
        "mode": "editable_text_failsafe",
        "slides": len(slides),
        "editable_slides": editable_slides,
        "raster_slides": 0,
        "errors": errors,
    }


def finalize_compile_html(
    storage_dir: str,
    *,
    project_id: str = "",
    run_id: str = "",
    include_notes: bool = True,
    fidelity_verify: bool | None = None,
) -> dict[str, Any]:
    """Job B: html_output/ → immutable downloadable PPTX (+ notes).

    Validation is an active remediation/observability loop, not a deck-level
    refusal switch. The compiler always preserves the authored HTML and native
    editability; it never swaps a complete slide for a generic or raster slide.
    If infrastructure cannot produce a new native artifact, the transaction
    returns the explicitly labelled previous immutable release when available.
    """
    html_dir = os.path.join(storage_dir, "html_output")
    os.makedirs(html_dir, exist_ok=True)
    from tasks.release_export import project_source_revision

    advisories: list[dict[str, Any]] = []

    def advise(stage: str, code: str, detail: Any, *, fallback: str = "") -> None:
        item: dict[str, Any] = {
            "stage": stage,
            "code": code,
            "detail": str(detail)[:4000],
        }
        if fallback:
            item["fallback"] = fallback
        advisories.append(item)
        print(f"[html_to_pptx] advisory {stage}:{code}: {str(detail)[:500]}")

    def previous_release_result(stage: str, detail: Any) -> dict[str, Any] | None:
        """Keep the deck task successful without publishing a broken rebuild."""
        from tasks.release_export import (
            ReleaseExportError,
            resolve_previous_released_pptx,
            resolve_released_pptx,
        )

        resolved = None
        for resolver in (resolve_previous_released_pptx, resolve_released_pptx):
            try:
                resolved = resolver(storage_dir)
                break
            except ReleaseExportError:
                continue
        if not resolved:
            return None
        previous_path, _previous_sha, previous_manifest = resolved
        advise(stage, "previous_immutable_release_served", detail, fallback="previous_release")
        return {
            "ok": True,
            "pptx_path": previous_path,
            "mode": "previous_immutable_release",
            "project_id": project_id,
            "release_manifest": previous_manifest,
            "speaker_notes": {"preserved_in_previous_release": True},
            "runtime_layout": {},
            "fidelity": {"skipped": True, "reason": stage},
            "advisories": advisories,
            "fallback": {
                "mode": "previous_immutable_release",
                "reason": str(detail)[:4000],
                "explicitly_labelled": True,
                "whole_slide_replaced": False,
            },
        }

    try:
        runtime_layout = validate_html_dir_runtime_layout(storage_dir, html_dir)
    except Exception as exc:  # noqa: BLE001
        runtime_layout = {
            "ok": False,
            "hard": [f"runtime_layout_unavailable: {exc}"],
            "slides": [],
        }
        advise(
            "runtime_layout",
            "validation_unavailable_nonblocking",
            exc,
            fallback="native_compile_preserving_authored_html",
        )
    if runtime_layout.get("hard"):
        print(f"[html_to_pptx] runtime layout failures={len(runtime_layout['hard'])}")
        details = "; ".join(str(item) for item in (runtime_layout.get("hard") or [])[:8])
        advise(
            "runtime_layout",
            "measured_findings_preserved_for_object_repair",
            details or "see validation/runtime_layout_qa.json",
            fallback="native_compile_preserving_authored_html",
        )
    # The validator may canonically repair Scene roots and refresh pipeline-
    # owned semantic hosts. Snapshot only after those deterministic writes;
    # the later comparison then detects real concurrent user edits rather than
    # rejecting the compiler's own normalization.
    source_revision = project_source_revision(storage_dir)
    exports = os.path.join(storage_dir, "exports")
    os.makedirs(exports, exist_ok=True)
    build_id = uuid.uuid4().hex
    build_dir = Path(exports) / f".build-{build_id}"
    build_dir.mkdir(parents=True, exist_ok=False)
    out = str(build_dir / "presentation.pptx")
    compile_mode = "html_deckhtml"
    fallback_meta: dict[str, Any] = {}
    try:
        from tasks.speaker_notes import resolve_slide_transition

        transition = resolve_slide_transition(storage_dir)
    except Exception:  # noqa: BLE001
        transition = None

    def _compile_once(*, relaxed: bool = False) -> None:
        compiled = compile_html_dir_to_pptx(
            html_dir,
            out,
            slide_transition=None if relaxed else transition,
            embed_fonts=not relaxed,
        )
        # Keep the staging contract even for injected/custom compilers that
        # return a different artifact path instead of honoring ``out``.
        staged = Path(out)
        compiled_path = Path(str(compiled)) if compiled else staged
        if not staged.is_file() and compiled_path.is_file():
            shutil.copy2(compiled_path, staged)
        if not staged.is_file():
            raise HtmlToPptxError("DeckHTML compile produced no staged PPTX")

    try:
        _compile_once()
    except Exception as native_error:  # noqa: BLE001
        advise(
            "deckhtml",
            "native_compile_failed",
            native_error,
            fallback="retry_native_without_font_embedding",
        )
        Path(out).unlink(missing_ok=True)
        try:
            _compile_once(relaxed=True)
            compile_mode = "html_deckhtml_relaxed"
            fallback_meta = {
                "mode": compile_mode,
                "reason": "primary_native_compile_failed",
            }
        except Exception as retry_error:  # noqa: BLE001
            shutil.rmtree(build_dir, ignore_errors=True)
            previous = previous_release_result(
                "deckhtml",
                f"primary={native_error}; relaxed={retry_error}",
            )
            if previous:
                return previous
            raise HtmlToPptxError(
                "Native editable PPTX compilation failed after bounded retry; "
                "no previous immutable release is available: "
                f"primary={native_error}; relaxed={retry_error}"
            ) from retry_error

    fidelity_meta: dict[str, Any] = {"skipped": True}
    run_verify = fidelity_verify is not False
    preview_dir = Path(storage_dir) / "preview"
    has_previews = preview_dir.is_dir() and any(preview_dir.glob("slide_*.png"))
    native_compiled = compile_mode.startswith("html_deckhtml")
    if run_verify and native_compiled and not has_previews:
        advise(
            "fidelity",
            "canonical_previews_unavailable_nonblocking",
            "No canonical preview PNGs were available",
            fallback="native_editable_output",
        )
    if run_verify and native_compiled and has_previews:
        try:
            from tasks.pptx_visual_verify import verify_pptx_against_previews

            fidelity_meta = verify_pptx_against_previews(storage_dir, pptx_path=out)
            print(
                f"[html_to_pptx] fidelity ok={fidelity_meta.get('ok')} "
                f"failures={len(fidelity_meta.get('failures') or [])}"
            )
            if not fidelity_meta.get("ok"):
                from tasks.layout_findings import native_fidelity_remediation_plan

                remediation = native_fidelity_remediation_plan(fidelity_meta)
                remediation_path = Path(storage_dir) / "validation" / "native_fidelity_remediation.json"
                remediation_path.parent.mkdir(parents=True, exist_ok=True)
                temporary = remediation_path.with_suffix(".json.tmp")
                temporary.write_text(
                    json.dumps(remediation, ensure_ascii=False, indent=2) + "\n",
                    encoding="utf-8",
                )
                os.replace(temporary, remediation_path)
                fidelity_meta["remediation"] = remediation
                failed_stems = ", ".join(
                    str(item.get("stem"))
                    for item in (fidelity_meta.get("failures") or [])[:4]
                    if isinstance(item, dict) and item.get("stem")
                )
                advise(
                    "fidelity",
                    "object_remediation_recorded_nonblocking",
                    (failed_stems or "unknown slides")
                    + "; object-first plan saved to validation/native_fidelity_remediation.json",
                    fallback="native_editable_output",
                )
                # Fidelity findings are an object-level QA ledger, not an
                # availability gate. The native candidate was compiled from
                # the current edited HTML and must be published as the current
                # release with advisories. Returning the previous immutable
                # release here makes the canonical export transaction reject
                # it as stale and leaves the user unable to download edits.
        except Exception as exc:  # noqa: BLE001
            advise(
                "fidelity",
                "verifier_unavailable_nonblocking",
                exc,
                fallback="native_editable_output",
            )

    notes_meta: dict[str, Any] = {"embedded": 0, "skipped": True, "reason": "disabled"}
    if include_notes:
        try:
            notes_meta = embed_speaker_notes_into_pptx(storage_dir, out, html_dir=html_dir)
            print(
                f"[html_to_pptx] speaker notes embedded={notes_meta.get('embedded')} "
                f"skipped={notes_meta.get('skipped')}"
            )
        except Exception as exc:  # noqa: BLE001
            notes_meta = {"embedded": 0, "error": str(exc), "preserved_source_notes": True}
            advise(
                "speaker_notes",
                "embedding_failed_nonblocking",
                exc,
                fallback="native_editable_output_with_source_notes_preserved",
            )

    live_source_revision = project_source_revision(storage_dir)
    if live_source_revision.get("sha256") != source_revision.get("sha256"):
        advise(
            "source_revision",
            "changed_during_compile",
            "Presentation sources changed during compilation",
            fallback="recompile_latest_native",
        )
        stabilized = False
        for _attempt in range(3):
            source_revision = project_source_revision(storage_dir)
            try:
                Path(out).unlink(missing_ok=True)
                try:
                    _compile_once()
                    compile_mode = "html_deckhtml"
                    fallback_meta = {}
                except Exception as native_rebuild_error:  # noqa: BLE001
                    advise(
                        "source_revision",
                        "latest_native_rebuild_retry",
                        native_rebuild_error,
                        fallback="retry_native_without_font_embedding",
                    )
                    Path(out).unlink(missing_ok=True)
                    _compile_once(relaxed=True)
                    compile_mode = "html_deckhtml_relaxed"
                    fallback_meta = {
                        "mode": compile_mode,
                        "reason": "primary_latest_native_rebuild_failed",
                    }
                if include_notes:
                    try:
                        notes_meta = embed_speaker_notes_into_pptx(
                            storage_dir, out, html_dir=html_dir
                        )
                    except Exception as exc:  # noqa: BLE001
                        notes_meta = {
                            "embedded": 0,
                            "error": str(exc),
                            "preserved_source_notes": True,
                        }
                        advise(
                            "speaker_notes",
                            "embedding_failed_after_rebuild_nonblocking",
                            exc,
                            fallback="native_editable_output_with_source_notes_preserved",
                        )
                if project_source_revision(storage_dir).get("sha256") == source_revision.get("sha256"):
                    stabilized = True
                    break
            except Exception as exc:  # noqa: BLE001
                shutil.rmtree(build_dir, ignore_errors=True)
                previous = previous_release_result("source_revision", exc)
                if previous:
                    return previous
                raise HtmlToPptxError(
                    "Latest-source native rebuild failed and no previous immutable "
                    f"release is available: {exc}"
                ) from exc
        if not stabilized:
            shutil.rmtree(build_dir, ignore_errors=True)
            previous = previous_release_result(
                "source_revision",
                "sources kept changing during bounded rebuild attempts",
            )
            if previous:
                return previous
            raise HtmlToPptxError(
                "Sources kept changing during bounded rebuild attempts and no previous "
                "immutable release is available"
            )

    try:
        source_id = str(source_revision.get("sha256") or "source")[:12]
        release_id = f"{source_id}-{build_id[:8]}"
        release_dir = Path(exports) / "releases" / release_id
        release_dir.mkdir(parents=True, exist_ok=True)
        released_pptx = release_dir / "presentation.pptx"
        released_shape_map = Path(f"{released_pptx}.shape-map.json")
        os.replace(out, released_pptx)
        staged_shape_map = Path(f"{out}.shape-map.json")
        if staged_shape_map.is_file():
            os.replace(staged_shape_map, released_shape_map)

        from tasks.release_export import write_release_manifest

        manifest = write_release_manifest(
            storage_dir,
            pptx_path=released_pptx,
            project_id=project_id,
            run_id=run_id,
            status="released_with_advisories" if advisories else "released",
            mode=compile_mode,
            source_revision=source_revision,
            release_id=release_id,
            advisories=advisories,
            compile_meta={
                "native_requested": True,
                "mode": compile_mode,
                "fallback": fallback_meta or None,
            },
        )
    except BaseException:
        shutil.rmtree(build_dir, ignore_errors=True)
        raise
    shutil.rmtree(build_dir, ignore_errors=True)
    return {
        "ok": True,
        "pptx_path": str(released_pptx),
        "mode": compile_mode,
        "project_id": project_id,
        "release_manifest": manifest,
        "speaker_notes": notes_meta,
        "runtime_layout": runtime_layout,
        "fidelity": fidelity_meta,
        "advisories": advisories,
        "fallback": fallback_meta,
    }
