"""Text measurement must come from the font file, and everyone must agree.

The old ratio table read every non-narrow character as 0.55em, so a condensed
display face and a wide geometric sans measured identically, and the fitter, the
capacity derivation and the QA checker each disagreed about whether a line fit.
"""

import os
import subprocess
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "sunairum_scripts",
))

from svg_to_pptx.drawingml.font_metrics import (  # noqa: E402
    coerce_weight,
    describe,
    measure_text,
    missing_glyphs,
    registry,
    resolve_face,
)
from svg_to_pptx.drawingml.utils import estimate_text_width  # noqa: E402
from tasks.layout_geometry import text_width_px  # noqa: E402


class FontDiscovery(unittest.TestCase):
    def test_the_shipped_pack_families_are_discovered(self):
        families = {name.lower() for name in registry(refresh=True).families}
        for expected in ("montserrat", "lato", "anton", "dm sans"):
            self.assertIn(expected, families, f"{expected} is shipped but not measurable")

    def test_a_weight_resolves_to_the_nearest_face(self):
        bold = resolve_face("Montserrat", 700)
        self.assertIsNotNone(bold)
        self.assertEqual(bold.weight, 700)
        # Montserrat ships no Thin, so the request must land on the lightest real face
        # rather than silently measuring a Bold.
        thin = resolve_face("Montserrat", 100)
        self.assertIsNotNone(thin)
        self.assertLessEqual(thin.weight, 500)

    def test_an_unknown_family_does_not_resolve(self):
        self.assertIsNone(resolve_face("Not A Real Font", 400))

    def test_css_weight_keywords_map_to_classes(self):
        self.assertEqual(coerce_weight("bold"), 700)
        self.assertEqual(coerce_weight("600"), 600)
        self.assertEqual(coerce_weight(None), 400)
        self.assertEqual(coerce_weight("semibold"), 600)


class MeasurementTruth(unittest.TestCase):
    def test_a_condensed_face_measures_narrower_than_a_geometric_sans(self):
        sample = "Attention Is All You Need"
        anton = measure_text(sample, 48.0, "Anton", 400)
        montserrat = measure_text(sample, 48.0, "Montserrat", 400)
        self.assertGreater(anton, 0)
        # This is the whole point of the change: the ratio table returned one
        # number for both, so a condensed hero title was measured as if it were
        # half again as wide as it renders.
        self.assertLess(anton, montserrat * 0.85)

    def test_width_scales_linearly_with_size(self):
        one = measure_text("Transformer", 20.0, "Montserrat", 400)
        two = measure_text("Transformer", 40.0, "Montserrat", 400)
        self.assertAlmostEqual(two, one * 2, places=4)

    def test_bold_is_wider_than_regular_in_the_same_family(self):
        regular = measure_text("Encoder Decoder", 32.0, "Montserrat", 400)
        bold = measure_text("Encoder Decoder", 32.0, "Montserrat", 700)
        self.assertGreater(bold, regular)

    def test_letter_spacing_is_added_between_clusters_only(self):
        plain = measure_text("ABC", 40.0, "Montserrat", 400)
        spaced = measure_text("ABC", 40.0, "Montserrat", 400, letter_spacing=5.0)
        self.assertAlmostEqual(spaced - plain, 10.0, places=4)

    def test_narrow_and_wide_characters_differ(self):
        narrow = measure_text("iii", 40.0, "Montserrat", 400)
        wide = measure_text("WWW", 40.0, "Montserrat", 400)
        self.assertGreater(wide, narrow * 1.5)


class AppAndCompilerAgree(unittest.TestCase):
    """One primitive, so a line the fitter accepts is not re-flagged by QA."""

    SAMPLES = (
        ("Scaled Dot-Product Attention", 32.0, "Montserrat", False),
        ("Multi-Head Attention", 44.0, "Anton", False),
        ("BLEU 28.4 on WMT 2014 EN-DE", 24.0, "Lato", True),
        ("Türkçe karakterler: ğüşıöç", 28.0, "Montserrat", False),
    )

    def test_the_app_and_the_compiler_return_the_same_width(self):
        for text, size, family, bold in self.SAMPLES:
            with self.subTest(text=text, family=family):
                app = text_width_px(text, size, 0.0, bold, family)
                compiler = estimate_text_width(
                    text, size, "700" if bold else "400", family,
                )
                self.assertAlmostEqual(app, compiler, places=6)

    def test_agreement_holds_without_a_family_too(self):
        app = text_width_px("fallback path", 30.0, 0.0, False, None)
        compiler = estimate_text_width("fallback path", 30.0, "400", None)
        self.assertAlmostEqual(app, compiler, places=6)


class GlyphCoverage(unittest.TestCase):
    def test_covered_text_reports_nothing_missing(self):
        self.assertEqual(missing_glyphs("Turkish: ğüşıöç ĞÜŞİÖÇ", "Montserrat", 400), [])

    def test_an_uncovered_glyph_is_reported(self):
        # A CJK ideograph is in none of the Latin pack faces, and a renderer would
        # substitute a different font with different metrics.
        absent = missing_glyphs("注意", "Montserrat", 400)
        self.assertEqual(absent, ["注", "意"])

    def test_coverage_of_an_unknown_family_is_not_asserted(self):
        self.assertEqual(missing_glyphs("anything", "Not A Real Font", 400), [])


class EmbeddingPermission(unittest.TestCase):
    def test_every_shipped_face_may_be_embedded(self):
        # W5b embeds the chosen faces into ppt/fonts so the recipient's PowerPoint
        # renders the metrics we measured. A face that forbids embedding would break
        # that silently.
        blocked = [
            entry["path"] for entry in describe()["faces"] if not entry["embeddable"]
        ]
        self.assertEqual(blocked, [])


class RendererAgreement(unittest.TestCase):
    """Calibrate against the renderer instead of trusting ourselves.

    Two components agreeing on the same wrong number is the failure this wave
    exists to end, so the reference is what LibreOffice actually draws.
    """

    @staticmethod
    def _soffice() -> str | None:
        for name in ("soffice", "libreoffice"):
            try:
                subprocess.run([name, "--version"], capture_output=True, timeout=60, check=True)
                return name
            except Exception:
                continue
        return None

    @staticmethod
    def _renderer_used(pdf: str, family: str) -> bool | None:
        """Whether the renderer really used ``family`` rather than substituting.

        Calibration is meaningless against a substituted font: a renderer that
        cannot find Montserrat draws Liberation Sans, which is about 11% narrower,
        and the mismatch says nothing about our metrics. Returns None when the
        question cannot be answered.
        """
        try:
            listing = subprocess.run(
                ["pdffonts", pdf], capture_output=True, text=True, timeout=60, check=True,
            ).stdout
        except Exception:
            return None
        compact = family.replace(" ", "").lower()
        return any(compact in line.replace(" ", "").lower() for line in listing.splitlines()[2:])

    def test_measured_width_matches_what_libreoffice_renders(self):
        soffice = self._soffice()
        if soffice is None:
            self.skipTest("LibreOffice is not available in this environment")

        import tempfile
        from pptx import Presentation
        from pptx.util import Emu, Pt

        text = "Attention Is All You Need"
        family, size = "Montserrat", 40.0
        expected = measure_text(text, size, family, 400)

        with tempfile.TemporaryDirectory() as work:
            deck = Presentation()
            deck.slide_width = Emu(int(1280 * 9525))
            deck.slide_height = Emu(int(720 * 9525))
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(int(1280 * 9525)), Emu(int(200 * 9525)))
            frame = box.text_frame
            frame.word_wrap = False
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.size = Pt(size * 0.75)  # px to pt
            run.font.name = family
            source = os.path.join(work, "probe.pptx")
            deck.save(source)

            try:
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "pdf", source, "--outdir", work],
                    capture_output=True, timeout=180, check=True,
                )
            except Exception as error:  # pragma: no cover - environment dependent
                self.skipTest(f"LibreOffice conversion unavailable: {error}")

            pdf = os.path.join(work, "probe.pdf")
            if not os.path.exists(pdf):
                self.skipTest("LibreOffice produced no PDF")

            used = self._renderer_used(pdf, family)
            if used is None:
                self.skipTest("pdffonts is unavailable, so substitution cannot be ruled out")
            if not used:
                self.skipTest(
                    f"the renderer substituted a different face for {family}; register the "
                    "shipped fonts with fontconfig before calibrating"
                )

            # pdftotext -bbox gives the laid-out glyph box, which is the renderer's
            # own answer to "how wide is this line".
            try:
                probe = subprocess.run(
                    ["pdftotext", "-bbox", pdf, "-"],
                    capture_output=True, text=True, timeout=120, check=True,
                )
            except Exception as error:  # pragma: no cover - environment dependent
                self.skipTest(f"pdftotext unavailable: {error}")

            import re
            xs = [
                (float(match.group(1)), float(match.group(2)))
                for match in re.finditer(r'xMin="([\d.]+)"[^>]*xMax="([\d.]+)"', probe.stdout)
            ]
            if not xs:
                self.skipTest("no text boxes recovered from the rendered PDF")
            # PDF user units are points; the deck is 1280px wide over 960pt.
            rendered_pt = max(x_max for _x_min, x_max in xs) - min(x_min for x_min, _x in xs)
            rendered_px = rendered_pt / 0.75

            self.assertGreater(rendered_px, 0)
            drift = abs(rendered_px - expected) / expected
            # A few percent covers hinting and the renderer's own rounding. The
            # ratio table used to be out by far more than this on Anton.
            self.assertLess(
                drift, 0.06,
                f"measured {expected:.1f}px but LibreOffice rendered {rendered_px:.1f}px",
            )


if __name__ == "__main__":
    unittest.main()
